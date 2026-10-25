"""Safe sparse linear expansion of local Excel aggregate formulas.

Retains the terminal amount cell, multiplicity and coefficient. Source terminal
amounts are the numeric endpoints; it never recalculates source prices/quantities
or external links. Supports only arithmetic, SUM, and exact integer TRUNC(...,0).
"""
from decimal import Decimal
from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import range_boundaries,get_column_letter

def numeric(v):
    if v is None or v=='':return Decimal(0)
    if isinstance(v,bool):raise ValueError('boolean aggregate value')
    x=Decimal(str(v).replace(',','').strip())
    if not x.is_finite():raise ValueError('nonfinite source number')
    return x
class Linear:
    def __init__(self,constant=0,terms=None):self.constant=Decimal(constant);self.terms=terms or {}
    def scaled(self,m):return Linear(self.constant*m,{k:v*m for k,v in self.terms.items()if v*m})
    def plus(self,other):
        terms=self.terms.copy()
        for k,v in other.terms.items():terms[k]=terms.get(k,Decimal(0))+v
        return Linear(self.constant+other.constant,{k:v for k,v in terms.items()if v})
    def value(self,sheets):return self.constant+sum(weight*numeric(sheets[sn][coord].get('value'))for(sn,coord),weight in self.terms.items())
class Expansion:
    def __init__(self,sheets,target,terminal_rows,price_col):
        self.sheets=sheets;self.target=target;self.rows=set(terminal_rows);self.cols={price_col+1,price_col+3,price_col+5,price_col+7};self.memo={}
    def cell(self,sn,coord,trail=()):
        if(sn,coord)in trail:raise ValueError('aggregate dependency cycle')
        if(sn,coord)in self.memo:return self.memo[(sn,coord)]
        c=self.sheets[sn].get(coord,{})
        if sn==self.target and c.get('row')in self.rows:
            if c.get('column',c.get('col'))in self.cols:return Linear(0,{(sn,coord):Decimal(1)})
            return Linear(numeric(c.get('value')))
        if c.get('formula'):
            value=Parser(c['formula'],self,sn,trail+((sn,coord),)).parse()
            if isinstance(value,list):raise ValueError('range outside aggregate function')
        else:value=Linear(numeric(c.get('value')))
        self.memo[(sn,coord)]=value
        return value
    def reference(self,value,sn,trail):
        ref=value.replace('$','')
        if '['in ref or '#REF!'in ref:raise ValueError('external/unresolved aggregate reference')
        if '!'in ref:
            sheet,ref=ref.rsplit('!',1);sheet=sheet.strip("'").replace("''","'")
        else:sheet=sn
        if sheet not in self.sheets:raise ValueError('missing aggregate sheet')
        l,t,r,b=range_boundaries(ref)
        values=[self.cell(sheet,f'{get_column_letter(col)}{row}',trail)for row in range(t,b+1)for col in range(l,r+1)]
        return values[0]if len(values)==1 else values
class Parser:
    def __init__(self,formula,expansion,sn,trail):
        self.ts=[t for t in Tokenizer(formula).items if t.type!='WHITE-SPACE'];self.i=0;self.exp=expansion;self.sn=sn;self.trail=trail
    def current(self):return self.ts[self.i]if self.i<len(self.ts)else None
    def parse(self):
        value=self.add()
        if self.i!=len(self.ts):raise ValueError('unsupported aggregate syntax')
        return value
    def add(self):
        a=self.mul()
        while self.current()and self.current().type=='OPERATOR-INFIX'and self.current().value in['+','-']:
            op=self.current().value;self.i+=1;b=self.mul();a=a.plus(b.scaled(Decimal(-1)if op=='-'else Decimal(1)))
        return a
    def mul(self):
        a=self.atom()
        while self.current()and self.current().type=='OPERATOR-INFIX'and self.current().value in['*','/']:
            op=self.current().value;self.i+=1;b=self.atom()
            if op=='/':
                if b.terms or b.constant==0:raise ValueError('nonlinear aggregate division')
                a=a.scaled(1/b.constant)
            elif a.terms and b.terms:raise ValueError('nonlinear aggregate multiplication')
            else:a=a.scaled(b.constant)if not b.terms else b.scaled(a.constant)
        return a
    def atom(self):
        token=self.current()
        if token is None:raise ValueError('empty aggregate operand')
        self.i+=1
        if token.type=='OPERATOR-PREFIX'and token.value in['+','-']:
            return self.atom().scaled(Decimal(-1)if token.value=='-'else Decimal(1))
        if token.type=='OPERAND':
            if token.subtype=='NUMBER':return Linear(numeric(token.value))
            if token.subtype=='RANGE':return self.exp.reference(token.value,self.sn,self.trail)
            raise ValueError('unsupported aggregate operand')
        if token.type=='PAREN'and token.subtype=='OPEN':
            x=self.add();end=self.current()
            if end is None or end.type!='PAREN'or end.subtype!='CLOSE':raise ValueError('aggregate parenthesis')
            self.i+=1;return x
        if token.type=='FUNC'and token.subtype=='OPEN':
            args=[];name=token.value[:-1].upper()
            while self.current()and not(self.current().type=='FUNC'and self.current().subtype=='CLOSE'):
                args.append(self.add())
                if self.current()and self.current().type=='SEP':self.i+=1
                elif self.current()and not(self.current().type=='FUNC'and self.current().subtype=='CLOSE'):raise ValueError('aggregate function separators')
            if self.current()is None:raise ValueError('unclosed aggregate function')
            self.i+=1
            if name=='SUM':
                value=Linear()
                for a in args:
                    for x in(a if isinstance(a,list)else[a]):value=value.plus(x)
                return value
            if name=='TRUNC'and len(args)==2 and not args[1].terms and args[1].constant==0:
                value=args[0];exact=value.value(self.exp.sheets)
                if exact==exact.to_integral_value() and value.constant==value.constant.to_integral_value() and all(w==w.to_integral_value()for w in value.terms.values()):return value
                raise ValueError('nonlinear truncation at aggregate')
            raise ValueError('unsupported aggregate function '+name)
        raise ValueError('unsupported aggregate token')

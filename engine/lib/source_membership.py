"""Source direct-cost membership from actual top-level formula dependencies.

Input cells accept {'row','col'/'column','value','formula'} as supplied by the
handoff extractor. This module never opens, writes, or evaluates external files.
"""
import re
import importlib.util
from pathlib import Path
from decimal import Decimal
from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import range_boundaries,get_column_letter
_spec=importlib.util.spec_from_file_location('independent_aggregate_linear',Path(__file__).with_name('aggregate_linear.py'))
_linear=importlib.util.module_from_spec(_spec);_spec.loader.exec_module(_linear)

def _normal(v):return re.sub(r'\s+','',str(v or ''))
def _number(v):
    if v is None or isinstance(v,bool):raise ValueError('missing aggregate number')
    n=Decimal(str(v).replace(',','').strip())
    if not n.is_finite():raise ValueError('nonfinite aggregate number')
    return n
def _layout(cells):
    for row in range(1,9):
        vals={c.get('column',c.get('col')):_normal(c.get('value')) for c in cells.values() if c.get('row')==row}
        found={label:[col for col,v in vals.items() if v==label] for label in ['품명','수량','단위']}
        if all(len(v)==1 for v in found.values()):
            nc,qc,uc=(found[label][0] for label in ['품명','수량','단위'])
            return {'name':nc,'price':max(qc,uc)+1,'total':max(qc,uc)+8,'scenario':max(qc,uc)+10}
    raise ValueError('unrecognized source detail layout')
def _refs(formula,sn):
    for t in Tokenizer(formula).items:
        if t.type!='OPERAND' or t.subtype!='RANGE':continue
        reference=t.value.replace('$','')
        if '[' in reference or '#REF!' in reference:raise ValueError('unresolved aggregate reference')
        if '!' in reference:
            sheet,reference=reference.rsplit('!',1);sheet=sheet.strip("'").replace("''", "'")
        else:sheet=sn
        l,t,r,b=range_boundaries(reference)
        for row in range(t,b+1):
            for col in range(l,r+1):yield sheet,f'{get_column_letter(col)}{row}'

def direct_membership(sheets,source_id,target_sheet,terminal_rows,price_col=None):
    """Return exact terminal rows and evidence; failed/missing totals stay UNRESOLVED.

    All referenced ancestors are traversed across local worksheets. Traversal
    stops at selected sheet terminal rows (including a one-row subtotal and a
    negative byproduct deduction). No union of unrelated excluded subtotals.
    """
    terminal_rows=set(terminal_rows);layout=_layout(sheets[target_sheet])
    pc=price_col or layout['price'];total_col=get_column_letter(pc+7)
    anchors=[]
    if target_sheet=='건축내역서':
        for sn in ['공종별집계표']:
            if sn not in sheets:continue
            for c in sheets[sn].values():
                col=c.get('column',c.get('col'))
                if col==1 and str(c.get('value') or '').strip().startswith('0101') and '건축공사' in str(c.get('value')):
                    anchors.append((sn,c['row'],'L',14))
    elif target_sheet=='기계내역서':
        sn='공종별집계표(기계)'
        if sn in sheets:
            for row in range(6,9):
                if sheets[sn].get(f'L{row}',{}).get('formula'):anchors.append((sn,row,'L',14))
    elif target_sheet=='토목내역서':
        for c in sheets[target_sheet].values():
            if c.get('column',c.get('col'))==layout['name'] and _normal(c.get('value'))=='순공사비합계':
                anchors.append((target_sheet,c['row'],total_col,pc+9))
    else:return {'status':'UNRESOLVED','terminal_rows':[],'anchor_evidence':[],'reason':'unsupported detail worksheet'}
    members=set();evidence=[]
    expansion=_linear.Expansion(sheets,target_sheet,terminal_rows,pc)
    for sn,row,col,gc in anchors:
        leaves=set()
        try:
            vectors={}
            source_pc=pc if sn==target_sheet else 5
            for key,offset in [('material',1),('labor',3),('expense',5)]:
                source_col=get_column_letter(source_pc+offset)
                vectors[key]=expansion.cell(sn,f'{source_col}{row}')
            for rr in terminal_rows:
                values={key:_number(sheets[target_sheet][f'{get_column_letter(pc+offset)}{rr}'].get('value'))for key,offset in [('material',1),('labor',3),('expense',5)]}
                weights={key:vectors[key].terms.get((target_sheet,f'{get_column_letter(pc+offset)}{rr}'),Decimal(0))for key,offset in [('material',1),('labor',3),('expense',5)]}
                required=[key for key,v in values.items()if v!=0]
                # A zero component reference cannot drag a helper row's unused
                # nonzero component into direct cost (e.g. building H84).
                if required:
                    if all(weights[key]==1 for key in required):leaves.add(rr)
                elif any(w==1 for w in weights.values()):leaves.add(rr)
            if not leaves:raise ValueError('no terminal source rows in aggregate')
            expected=_number(sheets[sn][f'{col}{row}'].get('value'))
            found=sum(_number(sheets[target_sheet][f'{total_col}{rr}'].get('value')) for rr in leaves)
            components={}
            for key,offset in [('material',1),('labor',3),('expense',5)]:
                subtotal=sum(_number(sheets[target_sheet][f'{get_column_letter(pc+offset)}{rr}'].get('value')) for rr in leaves)
                source_col=get_column_letter((pc if sn==target_sheet else 5)+offset)
                cached=_number(sheets[sn][f'{source_col}{row}'].get('value'))
                expanded=vectors[key].value(sheets)
                components[key]={'independent':str(subtotal),'formula_linear_expansion':str(expanded),'source_cached':str(cached),'status':'PASS' if subtotal==cached==expanded else 'FAIL'}
            ok=found==expected and all(x['status']=='PASS' for x in components.values())
            evidence.append({'source_id':source_id,'target_sheet':target_sheet,'anchor':f'{sn}!{col}{row}','scenario':sheets[sn].get(f'{get_column_letter(gc)}{row}',{}).get('value'),'terminal_count':len(leaves),'terminal_rows':sorted(leaves),'independent_total':str(found),'source_cached_total':str(expected),'component_checks':components,'status':'PASS' if ok else 'FAIL'})
            if ok:members|=leaves
        except Exception as exc:evidence.append({'source_id':source_id,'target_sheet':target_sheet,'anchor':f'{sn}!{col}{row}','status':'UNRESOLVED','reason':str(exc)})
    good=bool(evidence) and all(x['status']=='PASS' for x in evidence)
    return {'status':'VERIFIED_SOURCE_DIRECT_MEMBERSHIP' if good else 'UNRESOLVED','terminal_rows':sorted(members) if good else [],'anchor_evidence':evidence,'reason':None if good else 'top-level source aggregate unresolved'}

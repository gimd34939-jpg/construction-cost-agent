"""Reusable native Excel verification on new disposable copies only."""
from __future__ import annotations

import hashlib
import json
import re
import shutil
import subprocess
import uuid
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl

_CELL = re.compile(r"[A-Z]{1,3}[1-9][0-9]{0,6}")
MAX_NATIVE_COPY_PATH = 218


def _sha(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _read(path):
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


def _dump(path, data):
    Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def _errors(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    found = {}
    try:
        for ws in wb.worksheets:
            for row in ws:
                for cell in row:
                    if cell.data_type == "e":
                        found[f"{ws.title}!{cell.coordinate}"] = cell.value
    finally:
        wb.close()
    return found


def _matches(expected, actual):
    if expected is None:
        return actual is None, "0"
    if isinstance(expected, bool):
        return type(expected) is type(actual) and expected == actual, "0"
    try:
        wanted = Decimal(str(expected))
        if not wanted.is_finite() or isinstance(actual, bool) or actual is None:
            return False, "0"
        got = Decimal(str(actual))
        # Final whole-won amounts are exact. Fractional intermediates retain
        # Excel's binary floating-point tails; tolerance is below one won.
        tolerance = Decimal(0) if wanted == wanted.to_integral_value() else Decimal("0.000001")
        return got.is_finite() and abs(wanted - got) <= tolerance, str(tolerance)
    except InvalidOperation:
        return (actual == expected or (expected == "" and actual is None)), "0"


def _compare(checks, actual):
    found = {(c["sheet"], c["cell"]): c for c in actual}
    rows = []
    for check in checks:
        got = found.get((check["sheet"], check["cell"]))
        okay, tolerance = _matches(check["expected"], got["value"] if got else None)
        rows.append({**check, "actual": got["value"] if got else None, "formula": got.get("formula") if got else None,
                     "text": got.get("text") if got else None, "tolerance": tolerance, "pass": okay and got is not None})
    return rows


def verify_workbook(file, checks, output_dir, *, mutations=None, mode="AUTO"):
    """Verify check cells plus optional reversible mutation cases in a new copy.

    checks: [{sheet,cell,expected}]. mutations: [{name,inputs:[{sheet,cell,value}],
    checks:[{sheet,cell,expected}]}]. Source/result files are never saved by Excel.
    AUTO recalculates the whole new workbook only when it has no external links;
    linked workbooks recalculate the supplied dependency-ordered cells only.
    """
    file = Path(file).resolve()
    if file.suffix.lower() != ".xlsx" or not file.is_file():
        raise ValueError("Native 검증에는 실제 XLSX 파일이 필요합니다.")
    if not isinstance(checks, list) or not checks:
        raise ValueError("검증할 명시 셀 목록이 필요합니다.")
    for check in checks + [c for m in (mutations or []) for c in m.get("checks", [])]:
        if not isinstance(check.get("sheet"), str) or not check["sheet"] or not _CELL.fullmatch(check.get("cell", "")) or "expected" not in check:
            raise ValueError("검증 셀은 sheet/cell/expected를 명시해야 합니다.")
    if len({(c['sheet'], c['cell']) for c in checks}) != len(checks):
        raise ValueError("중복 검증 셀은 허용하지 않습니다.")
    for mutation in mutations or []:
        for value in mutation.get("inputs", []):
            if not _CELL.fullmatch(value.get("cell", "")) or not isinstance(value.get("sheet"), str):
                raise ValueError("경계시험 대상 셀 형식이 잘못되었습니다.")
            supplied = value.get("value")
            if isinstance(supplied, bool) or (isinstance(supplied, str) and supplied.startswith(("=", "+", "-", "@"))):
                raise ValueError("경계시험은 불리언 또는 수식 형태 문자열을 삽입하지 않습니다.")
    output = Path(output_dir).resolve()
    # Excel native Open still rejects deep paths. Keep disposable names short;
    # full source identity remains in the request/report and SHA256 fields.
    folder = output / ("n_" + uuid.uuid4().hex[:8])
    copy = folder / "v.xlsx"
    if len(str(copy)) > MAX_NATIVE_COPY_PATH:
        raise ValueError(
            f"Native Excel 검증 사본 경로가 운영상 한도 {MAX_NATIVE_COPY_PATH}자를 초과합니다 "
            f"({len(str(copy))}자). --output-dir를 더 짧은 경로로 지정하세요. "
            "원본은 파일시스템으로 읽어 짧은 검증 사본에 복사하므로 원본 경로를 바꿀 필요가 없습니다."
        )
    output.mkdir(parents=True, exist_ok=True)
    folder.mkdir()
    if copy == file or copy.exists():
        raise ValueError("검증 사본 경로가 원본과 같거나 이미 존재합니다.")
    before = _sha(file)
    original_errors = _errors(file)
    with zipfile.ZipFile(file) as z:
        external_links = sum(bool(re.fullmatch(r"xl/externalLinks/externalLink\d+\.xml", n)) for n in z.namelist())
    if mode == "AUTO":
        mode = "CHECKED_CELLS" if external_links else "WORKBOOK"
    if mode not in ("CHECKED_CELLS", "WORKBOOK") or (external_links and mode == "WORKBOOK"):
        raise ValueError("외부 링크가 있는 파일은 명시된 로컬 종속 셀만 재계산합니다.")
    shutil.copyfile(file, copy)
    request = {"source_path": str(file), "copy_path": str(copy), "checks": checks, "mutations": mutations or [],
               "mode": mode, "output": str(folder / "native_cells.json")}
    _dump(folder / "native_request.json", request)
    try:
        process = subprocess.run(["powershell.exe", "-NoLogo", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File",
                                  str(Path(__file__).with_name("verify_native_excel.ps1")), "-RequestPath", str(folder / "native_request.json")],
                                 capture_output=True, timeout=180)
    except subprocess.TimeoutExpired as error:
        after = _sha(file)
        (folder / "native_process_stdout.txt").write_bytes(error.stdout or b"")
        (folder / "native_process_stderr.txt").write_bytes(error.stderr or b"")
        report = {"status": "FAIL", "source_file": str(file), "source_unchanged": before == after,
                  "source_sha256_before": before, "source_sha256_after": after,
                  "reason": "Native Excel 검증이 180초 안에 완료되지 않았습니다. 검증 사본과 로그를 확인하세요.",
                  "folder": str(folder), "native_copy": str(copy), "native_copy_path_length": len(str(copy)),
                  "native_copy_path_limit": MAX_NATIVE_COPY_PATH}
        _dump(folder / "native_validation.json", report)
        return report
    (folder / "native_process_stdout.txt").write_bytes(process.stdout)
    (folder / "native_process_stderr.txt").write_bytes(process.stderr)
    after = _sha(file)
    if process.returncode or not Path(request["output"]).is_file():
        report = {"status": "FAIL", "source_file": str(file), "source_unchanged": before == after,
                  "source_sha256_before": before, "source_sha256_after": after,
                  "reason": "Native Excel 프로세스 실패", "returncode": process.returncode, "folder": str(folder)}
        _dump(folder / "native_validation.json", report)
        return report
    native = _read(request["output"])
    baseline = _compare(checks, native["baseline"])
    restored = _compare(checks, native["restored"])
    mutated = []
    for wanted, got in zip(mutations or [], native["mutations"]):
        compared = _compare(wanted["checks"], got["checks"])
        mutated.append({"name": wanted["name"], "inputs": wanted["inputs"], "checks": compared,
                        "pass": all(c["pass"] for c in compared)})
    final_errors = _errors(copy)
    new_errors = {k: v for k, v in final_errors.items() if k not in original_errors or original_errors[k] != v}
    passed = (before == after and all(c["pass"] for c in baseline + restored) and not new_errors
              and len(mutated) == len(mutations or []) and all(c["pass"] for c in mutated))
    report = {"status": "PASS" if passed else "FAIL", "source_file": str(file), "source_sha256_before": before,
              "source_sha256_after": after, "source_unchanged": before == after, "native_copy": str(copy),
              "native_copy_path_length": len(str(copy)), "native_copy_path_limit": MAX_NATIVE_COPY_PATH,
              "native_copy_sha256": _sha(copy), "folder": str(folder), "excel": {k: native[k] for k in ("version", "build", "mode", "update_links", "macros", "calculation")},
              "external_links_unrefreshed": external_links, "baseline_check_count": len(baseline), "baseline": baseline,
              "restored": restored, "mutations": mutated, "original_cached_errors": original_errors, "final_cached_errors": final_errors,
              "new_error_cells": new_errors,
              "scope": "명시 셀의 독립 기대값과 Excel 재계산 대조. 기존 오류·외부 링크를 해소한 검증 또는 전체 도급액 검증이 아님."}
    _dump(folder / "native_validation.json", report)
    return report


def verify_package(package, output_dir):
    """Verify comparison, edited bill, and generated calculation workbooks."""
    output = Path(output_dir)
    reports = {}
    if package.get("comparison_workbook"):
        original = _read(Path(package["folder"]) / "비교Excel검증.json")
        reports["comparison"] = verify_workbook(package["comparison_workbook"], original["checks"], output)
    edit = package.get("edited_copy") or {}
    if edit.get("workbook"):
        r = _read(edit["report"])
        checks = [{"sheet": x["sheet"], "cell": x["cell"], "expected": x["value"]} for x in r["recalculated_cells"]]
        reports["edited_bill"] = verify_workbook(edit["workbook"], checks, output)
    if edit.get("calculation_workbook"):
        path = Path(edit["calculation_workbook"])
        r = _read(path.parent / "Excel검증.json")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
        first_sheet = wb.sheetnames[0]
        wb.close()
        checks = [{"sheet": x.get("sheet", first_sheet), "cell": x["cell"], "expected": x["expected"]} for x in r["checks"]]
        reports["calculation"] = verify_workbook(path, checks, output)
    result = {"status": "PASS" if reports and all(r["status"] == "PASS" for r in reports.values()) else "FAIL",
              "reports": reports, "workbooks": len(reports), "checks": sum(r.get("baseline_check_count", 0) for r in reports.values()),
              "scope": "생성된 사본의 명시 계산 범위에 대한 native Excel 검증. 원본 오류/외부링크/전체 도급액 검증으로 확대하지 않음."}
    output.mkdir(parents=True, exist_ok=True)
    _dump(output / ("native_package_" + uuid.uuid4().hex[:8] + ".json"), result)
    return result


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--package")
    source.add_argument("--workbook")
    parser.add_argument("--checks-file", help="검증 셀 배열 또는 checks/mutations/mode가 있는 JSON")
    parser.add_argument("--mutations-file", help="선택적인 수량 변경·복원 시험 배열 JSON")
    parser.add_argument("--mode", choices=("AUTO", "CHECKED_CELLS", "WORKBOOK"), default=None)
    parser.add_argument("--output-dir", required=True)
    args = parser.parse_args()
    if args.package:
        if args.checks_file or args.mutations_file or args.mode:
            parser.error("--package는 checks/mutations/mode 개별 옵션과 함께 사용하지 않습니다.")
        result = verify_package(_read(args.package), args.output_dir)
        summary = {"status": result["status"], "workbooks": result["workbooks"], "checks": result["checks"]}
    else:
        if not args.checks_file:
            parser.error("--workbook에는 독립 기대값을 가진 --checks-file이 필요합니다.")
        payload = _read(args.checks_file)
        checks = payload if isinstance(payload, list) else payload["checks"]
        mutations = _read(args.mutations_file) if args.mutations_file else (payload.get("mutations", []) if isinstance(payload, dict) else [])
        mode = args.mode or (payload.get("mode", "AUTO") if isinstance(payload, dict) else "AUTO")
        result = verify_workbook(args.workbook, checks, args.output_dir, mutations=mutations, mode=mode)
        summary = {"status": result["status"], "checks": result.get("baseline_check_count", 0),
                   "mutations": len(result.get("mutations", [])), "folder": result.get("folder"),
                   "source_unchanged": result["source_unchanged"]}
    print(json.dumps(summary, ensure_ascii=False))
    raise SystemExit(0 if result["status"] == "PASS" else 1)

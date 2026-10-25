"""Preserve an immutable XLSX and edit a new OOXML copy only.

Supported: explicit numeric contract leaf updates, and one-row additions beside
an existing additive contract leaf. Complex target drawings/tables/array formulas
and unresolved dependent calculations are blocked rather than guessed.
"""
from pathlib import Path
from decimal import Decimal,localcontext
from collections import defaultdict,deque
import copy,datetime,hashlib,importlib.util,io,json,posixpath,re,zipfile
from lxml import etree as ET
from openpyxl.formula import Tokenizer
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import range_boundaries,get_column_letter,coordinate_from_string,column_index_from_string

OUT=Path(__file__).resolve().parents[1]
WORKSPACE=OUT.parents[1]
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
N={'s':NS}
class EditError(ValueError):pass
def _sha(p):
    with p.open('rb')as f:return hashlib.file_digest(f,'sha256').hexdigest()
def _number(v,label='number'):
    if v is None or isinstance(v,bool)or not re.fullmatch(r'-?(?:\d+(?:\.\d*)?|\.\d+)',str(v)):
        raise EditError('명시된 유한 십진수 필요: '+label)
    x=Decimal(str(v))
    if not x.is_finite()or abs(x)>Decimal('1000000000000000'):raise EditError('숫자 범위 초과: '+label)
    return x
def _normal(v):return re.sub(r'\s+','',str(v or ''))
def _q(tag):return '{'+NS+'}'+tag
def _xml(data):return ET.fromstring(data,parser=ET.XMLParser(resolve_entities=False,no_network=True,remove_blank_text=False))
def _serialize(node):return ET.tostring(node,xml_declaration=True,encoding='UTF-8',standalone=True)
def _module(name,p):
    spec=importlib.util.spec_from_file_location(name,p);m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m);return m
_ast=None
def _evaluator():
    global _ast
    if _ast is None:_ast=_module('source_edit_decimal_evaluator',OUT/'engine/lib/reference_ast.py')
    return _ast

class Book:
    def __init__(self,path):
        with zipfile.ZipFile(path)as z:
            self.infos=z.infolist();self.parts={i.filename:z.read(i.filename)for i in self.infos}
        if any(n.startswith('_xmlsignatures/')for n in self.parts):raise EditError('전자서명 문서는 자동 편집 지원하지 않습니다')
        self.workbook=_xml(self.parts['xl/workbook.xml']);rels=_xml(self.parts['xl/_rels/workbook.xml.rels'])
        relationships={r.get('Id'):r.get('Target')for r in rels}
        self.paths={};self.ids={}
        for sheet in self.workbook.find(_q('sheets')):
            p=relationships[sheet.get('{'+REL+'}id')]
            self.paths[sheet.get('name')]=posixpath.normpath(posixpath.join('xl',p))if not p.startswith('/')else p.lstrip('/')
            self.ids[sheet.get('name')]=int(sheet.get('sheetId'))
        self.trees={sn:_xml(self.parts[p])for sn,p in self.paths.items()}
        self.strings=[]
        if 'xl/sharedStrings.xml'in self.parts:
            for si in _xml(self.parts['xl/sharedStrings.xml']).iter(_q('si')):
                self.strings.append(''.join(t.text or''for t in si.iter(_q('t'))))
        self.changed=set();self.cells={};self.formulas={};self._index()
    def _index(self):
        self.cells={};self.formulas={}
        for sn,tree in self.trees.items():
            cells={c.get('r'):c for c in tree.iter(_q('c'))};self.cells[sn]=cells
            masters={}
            for addr,c in cells.items():
                f=c.find(_q('f'))
                if f is not None and f.get('t')=='shared'and f.text:masters[f.get('si')]=(addr,f.text)
            fs={}
            for addr,c in cells.items():
                f=c.find(_q('f'))
                if f is None:continue
                if f.text:fs[addr]='='+f.text
                elif f.get('t')=='shared'and f.get('si')in masters:
                    origin,text=masters[f.get('si')];fs[addr]=Translator('='+text,origin=origin).translate_formula(addr)
                else:raise EditError('해석할 수 없는 공유 수식: '+sn+'!'+addr)
            self.formulas[sn]=fs
    def value(self,c):
        if c is None:return None
        v=c.find(_q('v'));typ=c.get('t')
        if typ=='s':return self.strings[int(v.text)]if v is not None else None
        if typ=='inlineStr':return ''.join(t.text or''for t in c.iter(_q('t')))
        if v is None or v.text is None:return None
        if typ in['str','e']:return v.text
        if typ=='b':return v.text=='1'
        try:return Decimal(v.text)
        except Exception:return v.text
    def model(self):
        return {sn:{addr:{'value':self.value(c),'formula':self.formulas[sn].get(addr),'error':c.get('t')=='e','row':coordinate_from_string(addr)[1],'col':column_index_from_string(coordinate_from_string(addr)[0])}for addr,c in cells.items()}for sn,cells in self.cells.items()}
    def layout(self,sn):
        cells=self.model()[sn]
        for row in range(1,9):
            vals={c['col']:_normal(c['value'])for c in cells.values()if c['row']==row}
            hits={label:[col for col,v in vals.items()if v==label]for label in['품명','수량','단위']}
            if all(len(v)==1 for v in hits.values()):
                nc,qc,uc=(hits[label][0]for label in['품명','수량','단위'])
                return {'name':nc,'specification':nc+1,'quantity':qc,'unit':uc,'price':max(qc,uc)+1,'scenario':max(qc,uc)+10}
        raise EditError('지원되지 않는 내역서 머리글: '+sn)
    def leaf_rows(self,sn,layout):
        out=set()
        for addr,c in self.cells[sn].items():
            col,row=coordinate_from_string(addr)
            if column_index_from_string(col)!=layout['name']:continue
            name=self.value(c);unit=self.value(self.cells[sn].get(f'{get_column_letter(layout["unit"])}{row}'))
            if isinstance(name,str)and str(unit or'').strip()and _normal(unit)not in['단위','수량','규격']:out.add(row)
        return out
    def set_cell(self,sn,address,value,formula=False):
        c=self.cells[sn].get(address)
        if c is None:
            row=coordinate_from_string(address)[1];data=self.trees[sn].find(_q('sheetData'))
            rr=next((r for r in data if int(r.get('r'))==row),None)
            if rr is None:raise EditError('기존 내역 행을 찾을 수 없습니다')
            c=ET.Element(_q('c'),r=address);rr.append(c)
            rr[:]=sorted(rr,key=lambda x:column_index_from_string(coordinate_from_string(x.get('r'))[0]))
            self.cells[sn][address]=c
        for child in list(c):
            if child.tag in[_q('f'),_q('v'),_q('is')]:c.remove(child)
        c.attrib.pop('t',None)
        if formula:
            ET.SubElement(c,_q('f')).text=str(value).lstrip('=');self.formulas[sn][address]='='+str(value).lstrip('=')
        elif isinstance(value,Decimal):ET.SubElement(c,_q('v')).text=format(value,'f');self.formulas[sn].pop(address,None)
        else:
            c.set('t','inlineStr');inline=ET.SubElement(c,_q('is'));t=ET.SubElement(inline,_q('t'));t.text=str(value);t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');self.formulas[sn].pop(address,None)
        self.changed.add(self.paths[sn])
    def normalize_formula(self,sn,address,formula):
        c=self.cells[sn][address];f=c.find(_q('f'))
        if f is None:f=ET.SubElement(c,_q('f'))
        f.attrib.clear();f.text=formula.lstrip('=');self.formulas[sn][address]=formula
        self.changed.add(self.paths[sn])

def _shift_address(ref,insertion):
    def one(match):
        col,absolute,row=match.groups();r=int(row)
        if r>=insertion:
            if r>=1048576:raise EditError('Excel 최대 행을 넘는 삽입')
            r+=1
        return col+absolute+str(r)
    return re.sub(r'(\$?[A-Z]{1,3})(\$?)(\d+)',one,ref)
def _shift_formula(formula,context,target,insertion):
    items=Tokenizer(formula).items;out=[]
    for t in items:
        v=t.value
        if t.type=='OPERAND'and t.subtype=='RANGE':
            if '!'in v:prefix,addr=v.rsplit('!',1);sheet=prefix.strip("'").replace("''","'")
            else:prefix=None;addr=v;sheet=context
            if sheet==target and'['not in str(prefix or''):
                if re.fullmatch(r'\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?',addr):addr=_shift_address(addr,insertion)
                elif re.fullmatch(r'\$?\d+:\$?\d+',addr):
                    addr=re.sub(r'\d+',lambda m:str(int(m[0])+1 if int(m[0])>=insertion else int(m[0])),addr)
                v=(prefix+'!'if prefix is not None else'')+addr
        out.append(v)
    return '='+''.join(out)
def _local_refs(formula,sn,book):
    for t in Tokenizer(formula).items:
        if t.type!='OPERAND'or t.subtype!='RANGE':continue
        ref=t.value.replace('$','')
        if '['in ref or'#REF!'in ref:continue
        if '!'in ref:sheet,addr=ref.rsplit('!',1);sheet=sheet.strip("'").replace("''","'")
        else:sheet,addr=sn,ref
        if sheet not in book.trees:continue
        if not re.fullmatch(r'[A-Z]+\d+(?::[A-Z]+\d+)?',addr):continue
        l,t,r,b=range_boundaries(addr)
        if(r-l+1)*(b-t+1)>200000:raise EditError('지원 범위를 넘는 원문 수식 참조')
        for row in range(t,b+1):
            for col in range(l,r+1):yield sheet,f'{get_column_letter(col)}{row}'

def _insert(book,sn,row,operation,layout):
    insertion=row+1;path=book.paths[sn]
    relpath=posixpath.join(posixpath.dirname(path),'_rels',posixpath.basename(path)+'.rels')
    if relpath in book.parts:
        for rel in _xml(book.parts[relpath]):
            if rel.get('Type','').split('/')[-1]in['drawing','table','vmlDrawing','comments','threadedComment']:
                raise EditError('그림/표/주석 위치가 있는 대상 시트의 행 삽입은 지원하지 않습니다')
    if any(f.get('t')=='array'for f in book.trees[sn].iter(_q('f'))):raise EditError('대상 시트 배열 수식의 행 삽입은 지원하지 않습니다')
    if book.trees[sn].find(_q('tableParts'))is not None:raise EditError('표 범위의 행 삽입은 지원하지 않습니다')
    if any(p.startswith(('xl/charts/','xl/pivotTables/','xl/pivotCache'))for p in book.parts):raise EditError('차트/피벗 참조가 있는 문서의 행 삽입은 지원하지 않습니다')
    for merged in book.trees[sn].iter(_q('mergeCell')):
        l,t,r,b=range_boundaries(merged.get('ref'))
        if t<insertion<=b:raise EditError('세로 병합 범위 내부의 행 삽입은 지원하지 않습니다')
    pc=layout['price'];leaves=book.leaf_rows(sn,layout);model=book.model()
    linear=_module('edit_parent_linear',OUT/'engine/lib/aggregate_linear.py')
    expansion=linear.Expansion(model,sn,leaves,pc)
    # Identify the first local additive subtotal formulas that list this leaf.
    parents=[]
    for addr,formula in book.formulas[sn].items():
        col,rr=coordinate_from_string(addr);cc=column_index_from_string(col)
        if rr in leaves or cc not in[pc+1,pc+3,pc+5,pc+7]:continue
        references=list(_local_refs(formula,sn,book));leafaddr=f'{col}{row}'
        if(sn,leafaddr)not in references:continue
        vector=expansion.cell(sn,addr)
        if vector.terms.get((sn,leafaddr),Decimal(0))!=1:raise EditError('신규 행을 연결할 원문 소계가 단순 가산식이 아닙니다: '+addr)
        covered=False
        for t in Tokenizer(formula).items:
            if t.type=='OPERAND'and t.subtype=='RANGE'and'!'not in t.value and':'in t.value:
                ref=t.value.replace('$','');l,top,r,b=range_boundaries(ref)
                if l<=cc<=r and top<=row and b>=insertion:covered=True
        parents.append((addr,formula,col,covered))
    if not parents:raise EditError('원문 소계에 연결할 직접 합계 수식을 찾지 못했습니다')
    # Preserve formula semantics before any cell-coordinate moves.
    oldformulas={s:f.copy()for s,f in book.formulas.items()}
    data=book.trees[sn].find(_q('sheetData'));original=next((r for r in data if int(r.get('r'))==row),None)
    if original is None:raise EditError('복제할 원문 행이 없습니다')
    clone=copy.deepcopy(original)
    for rr in reversed(list(data)):
        current=int(rr.get('r'))
        if current>=insertion:
            rr.set('r',str(current+1))
            for c in rr:c.set('r',_shift_address(c.get('r'),insertion))
    clone.set('r',str(insertion))
    for c in clone:
        col,_=coordinate_from_string(c.get('r'));c.set('r',f'{col}{insertion}')
    original_index=list(data).index(original);data.insert(original_index+1,clone)
    # Shift coordinate-bearing worksheet features and all local cross-sheet refs.
    for s,tree in book.trees.items():
        if s==sn:
            for node in tree.iter():
                if node.tag==_q('c')or node.tag==_q('row'):continue
                for attr in['ref','sqref','activeCell']:
                    if attr in node.attrib:
                        old=node.get(attr);new=_shift_address(old,insertion)
                        if new!=old:node.set(attr,new);book.changed.add(book.paths[s])
                if node.tag==_q('brk')and node.get('id')and int(node.get('id'))>=insertion:
                    node.set('id',str(int(node.get('id'))+1));book.changed.add(book.paths[s])
        for tag in['formula','formula1','formula2']:
            for node in tree.iter(_q(tag)):
                if node.text:
                    shifted=_shift_formula('='+node.text,s,sn,insertion)
                    if shifted!='='+node.text:node.text=shifted[1:];book.changed.add(book.paths[s])
        for hyperlink in tree.iter(_q('hyperlink')):
            if hyperlink.get('location'):
                original_location=hyperlink.get('location');shifted=_shift_formula('='+original_location,s,sn,insertion)[1:]
                if shifted!=original_location:hyperlink.set('location',shifted);book.changed.add(book.paths[s])
    book.changed.add(path);book._index()
    for s,fs in oldformulas.items():
        for oldaddr,f in fs.items():
            newaddr=_shift_address(oldaddr,insertion)if s==sn else oldaddr
            shifted=_shift_formula(f,s,sn,insertion)
            if shifted!=f or(s==sn and newaddr!=oldaddr):book.normalize_formula(s,newaddr,shifted)
    # Shared blocks in modified worksheets are made explicit; all formulas remain
    # semantically identical, preventing followers referencing an edited master.
    for s in book.trees:
        if book.paths[s]in book.changed:
            for address,formula in book.formulas[s].copy().items():book.normalize_formula(s,address,formula)
    for addr,f,col,covered in parents:
        newaddr=_shift_address(addr,insertion);shifted=_shift_formula(f,sn,sn,insertion)
        if not covered:shifted=shifted+f'+{col}{insertion}'
        book.normalize_formula(sn,newaddr,shifted)
    for name in book.workbook.iter(_q('definedName')):
        if name.text:
            context=list(book.paths)[int(name.get('localSheetId'))]if name.get('localSheetId')is not None else''
            shifted=_shift_formula('='+name.text,context,sn,insertion)[1:]
            if shifted!=name.text:name.text=shifted;book.changed.add('xl/workbook.xml')
    return insertion,[{'parent_before':a,'parent_after':_shift_address(a,insertion),'range_already_expands':covered}for a,f,c,covered in parents]

def _apply_row(book,sn,row,op,layout,added=False):
    q=_number(op['quantity'],'quantity');parts={k:_number(op['unit_components'][k],k)for k in['material','labor','expense']}
    if any(p<0 for p in parts.values()):raise EditError('공종 단가는 음수일 수 없습니다')
    scale=_number(op.get('quantity_scale','1'),'quantity_scale')
    if scale not in[Decimal(1),Decimal('0.01')]:raise EditError('지원하지 않는 수량 환산계수')
    changed=set();qc=get_column_letter(layout['quantity']);pc=layout['price']
    book.set_cell(sn,f'{qc}{row}',q);changed.add((sn,f'{qc}{row}'))
    for j,key in enumerate(['material','labor','expense']):
        addr=f'{get_column_letter(pc+j*2)}{row}';book.set_cell(sn,addr,parts[key]);changed.add((sn,addr))
    if added:
        for key,col in[('name',layout['name']),('specification',layout['specification']),('unit',layout['unit'])]:book.set_cell(sn,f'{get_column_letter(col)}{row}',str(op.get(key,'')))
        for j in range(3):
            price=get_column_letter(pc+j*2);amount=get_column_letter(pc+j*2+1)
            addr=f'{amount}{row}';book.set_cell(sn,addr,f'=TRUNC({qc}{row}*{price}{row}*{format(scale,"f")},0)',True);changed.add((sn,addr))
        total=get_column_letter(pc+6);amounttotal=get_column_letter(pc+7)
        book.set_cell(sn,f'{total}{row}',f'=SUM({get_column_letter(pc)}{row},{get_column_letter(pc+2)}{row},{get_column_letter(pc+4)}{row})',True)
        book.set_cell(sn,f'{amounttotal}{row}',f'=SUM({get_column_letter(pc+1)}{row},{get_column_letter(pc+3)}{row},{get_column_letter(pc+5)}{row})',True)
        changed|={(sn,f'{total}{row}'),(sn,f'{amounttotal}{row}')}
    return changed
def _recalculate(book,changed):
    model=book.model();deps={};reverse=defaultdict(set)
    for sn,fs in book.formulas.items():
        for addr,f in fs.items():
            node=(sn,addr);ds=set(_local_refs(f,sn,book));deps[node]=ds
            for d in ds:reverse[d].add(node)
    closure=set(changed);queue=deque(changed)
    while queue:
        for child in reverse[queue.popleft()]:
            if child not in closure:closure.add(child);queue.append(child)
    formula_nodes={n for n in closure if n in deps}
    degrees={n:len(deps[n]&formula_nodes)for n in formula_nodes};ready=deque(n for n,d in degrees.items()if d==0)
    updated=[];evaluator=_evaluator()
    with localcontext()as ctx:
        ctx.prec=50
        while ready:
            sn,addr=ready.popleft();cell=model[sn][addr];f=cell['formula']
            try:value=evaluator.Evaluator(model,sn,recursive=False).evaluate(evaluator.AST(f,model.keys()).parse())
            except Exception as exc:raise EditError('변경 의존 수식을 확정 계산할 수 없습니다: '+sn+'!'+addr+' / '+str(exc))from exc
            if isinstance(value,list)or value is None:raise EditError('확정되지 않은 의존 수식 결과: '+sn+'!'+addr)
            c=book.cells[sn][addr];v=c.find(_q('v'))
            if v is None:v=ET.SubElement(c,_q('v'))
            if isinstance(value,bool):c.set('t','b');v.text='1'if value else'0'
            elif isinstance(value,(Decimal,int,float)):
                number=Decimal(str(value))
                if not number.is_finite():raise EditError('비정상 수식 숫자')
                c.attrib.pop('t',None);v.text=format(number,'f')
            elif isinstance(value,str):c.set('t','str');v.text=value
            else:raise EditError('지원되지 않는 수식 결과')
            cell['value']=value;cell['error']=False;book.changed.add(book.paths[sn]);updated.append({'sheet':sn,'cell':addr,'formula':f,'value':str(value)})
            for child in reverse[(sn,addr)]:
                if child in degrees:
                    degrees[child]-=1
                    if degrees[child]==0:ready.append(child)
    if len(updated)!=len(formula_nodes):raise EditError('변경 수식에 순환 참조가 있습니다')
    return updated

def apply_edit(source_path,expected_sha256,operations,output_path):
    source=Path(source_path).resolve();dest=Path(output_path).resolve()
    if source==dest or dest.exists():raise EditError('원본/기존 결과는 덮어쓸 수 없습니다')
    if source.suffix.lower()!='.xlsx'or dest.suffix.lower()!='.xlsx':raise EditError('XLSX 사본 편집만 지원합니다')
    if not source.is_file()or _sha(source)!=expected_sha256:raise EditError('원본 SHA256이 요청 근거와 다릅니다')
    if not isinstance(operations,list)or not operations:raise EditError('편집할 작업이 필요합니다')
    book=Book(source);changed=set();edits=[];insertions=[];prior_insertions=[]
    for op in operations:
        if not isinstance(op,dict)or op.get('operation')not in['update','add_after']:raise EditError('지원되지 않는 편집 작업')
        if any(op.get(k)is None for k in['sheet','row','quantity','unit_components','record_id','evidence']):raise EditError('편집 근거/필수 입력이 누락됐습니다')
        sn=op['sheet'];requested_row=op['row']
        if sn not in book.trees or type(requested_row)is not int or requested_row<=0:raise EditError('대상 시트/행을 확인하세요')
        row=requested_row+sum(s==sn and after<requested_row for s,after in prior_insertions)
        if sn not in['건축내역서','기계내역서','토목내역서']:raise EditError('지정 내역서 말단행만 편집할 수 있습니다')
        layout=book.layout(sn)
        if row not in book.leaf_rows(sn,layout):raise EditError('원문 공종 말단행을 찾을 수 없습니다')
        if set(op['unit_components'])!={'material','labor','expense'}:raise EditError('3비목 단가가 모두 필요합니다')
        added=op['operation']=='add_after';originalrow=row
        if added:
            if not str(op.get('name','')).strip()or not str(op.get('unit','')).strip():raise EditError('신규 공종명/단위가 필요합니다')
            row,parents=_insert(book,sn,row,op,layout);insertions.append({'sheet':sn,'after_row':originalrow,'new_row':row,'parents':parents})
            prior_insertions.append((sn,requested_row))
            # Coordinates in previously changed rows/dependency sources shift too.
            changed={(s,_shift_address(a,row)if s==sn else a)for s,a in changed}
            changed|={(sn,p['parent_after'])for p in parents}
        changed|=_apply_row(book,sn,row,op,layout,added)
        edits.append({'operation':op['operation'],'sheet':sn,'requested_row':requested_row,'output_row':row,'record_id':op['record_id'],'evidence':op['evidence']})
    recalculated=_recalculate(book,changed)
    # A rebuilt dependency chain is safer than a stale pre-insertion chain.
    removed=set()
    if 'xl/calcChain.xml'in book.parts:
        removed.add('xl/calcChain.xml')
        rels=_xml(book.parts['xl/_rels/workbook.xml.rels'])
        for rel in list(rels):
            if rel.get('Type','').endswith('/calcChain'):rels.remove(rel)
        book.parts['xl/_rels/workbook.xml.rels']=_serialize(rels);book.changed.add('xl/_rels/workbook.xml.rels')
        ct=_xml(book.parts['[Content_Types].xml'])
        for node in list(ct):
            if node.get('PartName')=='/xl/calcChain.xml':ct.remove(node)
        book.parts['[Content_Types].xml']=_serialize(ct);book.changed.add('[Content_Types].xml')
    calc=book.workbook.find(_q('calcPr'))
    if calc is None:calc=ET.SubElement(book.workbook,_q('calcPr'))
    calc.set('fullCalcOnLoad','1');calc.set('forceFullCalc','1');calc.set('calcMode','auto');book.changed.add('xl/workbook.xml')
    for sn,p in book.paths.items():
        if p in book.changed:
            for address,formula in book.formulas[sn].copy().items():book.normalize_formula(sn,address,formula)
            book.parts[p]=_serialize(book.trees[sn])
    book.parts['xl/workbook.xml']=_serialize(book.workbook)
    stream=io.BytesIO()
    with zipfile.ZipFile(stream,'w')as z:
        for info in book.infos:
            if info.filename not in removed:z.writestr(info,book.parts[info.filename])
    if _sha(source)!=expected_sha256:raise EditError('처리 중 원본이 변경됐습니다. 결과를 저장하지 않습니다')
    dest.parent.mkdir(parents=True,exist_ok=True)
    with dest.open('xb')as f:f.write(stream.getvalue())
    after=_sha(source)
    if after!=expected_sha256:raise EditError('원본 보존 검증 실패')
    return {'status':'EDITED_SOURCE_COPY','source_path':str(source),'source_sha256_before':expected_sha256,'source_sha256_after':after,'source_unchanged':True,'output_path':str(dest),'output_sha256':_sha(dest),'operations':edits,'insertions':insertions,'recalculated_local_cells':len(recalculated),'recalculated_cells':recalculated,'changed_ooxml_parts':sorted(book.changed),'removed_parts':sorted(removed),'preserved_parts':sorted(set(book.parts)-book.changed-removed),'scope':'원문 계약판본의 지정 말단행·명시된 가산 소계와 변경 의존 수식. 현재 견적/자동 제경비 확정 아님.','source_scope_warnings':['선택 원문 구간만 수정; 다른 연차/시나리오 물량 배분은 별도 명시 필요','현재 원본 총괄과 연차 간 기존 차이 자동 교정 안 함'],'created_utc':datetime.datetime.now(datetime.timezone.utc).isoformat()}

"""Source-linked multi-project composition, quantity takeoff and price comparison."""
from pathlib import Path
from decimal import Decimal
import datetime,hashlib,json,sqlite3
import cost_engine as cost
OUT=Path(__file__).resolve().parents[1]
RULE_DB=OUT/'data/공사구성_규칙.db'
class AssemblyError(ValueError):pass
def norm(x):return ''.join(str(x).lower().split())
def sha(p):
 with Path(p).open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()
def rules():
 c=sqlite3.connect(RULE_DB.as_uri()+'?mode=ro',uri=True);c.row_factory=sqlite3.Row;c.execute('PRAGMA query_only=ON');return c
def unit(x):
 return {'M':'m','M2':'m2','㎡':'m2','M3':'m3','㎥':'m3','EA':'count','개':'count','개소':'count','매':'count','SET':'count','set':'count','조':'count','Ton':'t','TON':'t','ton':'t','톤':'t'}.get(str(x),cost.canonical_unit(x))
def verify_file(evidence):
 if not isinstance(evidence,dict):raise AssemblyError('근거 객체가 필요합니다')
 p=Path(evidence.get('path',evidence.get('source_path','')))
 expected=evidence.get('sha256',evidence.get('source_sha256'))
 location=evidence.get('location',evidence.get('source_location'))
 if not p.is_absolute() or not p.is_file() or not expected or not isinstance(location,str) or not location.strip() or sha(p)!=expected:raise AssemblyError('근거 절대경로·SHA·문자열 위치를 확인하세요')
 return {'path':str(p),'sha256':expected,'location':location}
def list_types(text=''):
 with rules() as c:
  records=[json.loads(r[0]) for r in c.execute('SELECT payload_json FROM assembly_types ORDER BY type_id')]
 return [{'type_id':r['type_id'],'name':r['name'],'discipline':r['discipline'],'aliases':r.get('aliases',[]),'components':len(r['components']),'project_presets':[p['project_id'] for p in r.get('presets',[])]} for r in records if not text or any(norm(text) in norm(x) for x in [r['name'],r['type_id'],*r.get('aliases',[])])]
def project_list():
 with rules() as c:return [json.loads(r[0]) for r in c.execute('SELECT payload_json FROM projects')]
def external_item(record_id):
 with rules() as c:r=c.execute('SELECT payload_json FROM external_contracts WHERE record_id=?',(record_id,)).fetchone()
 return json.loads(r[0]) if r else None
def get_record(record_id):
 if str(record_id).startswith('EXT:'):
  v=external_item(record_id)
  if v is None:raise AssemblyError('외부 계약 행을 찾을 수 없습니다')
  return v
 with cost.connect() as c:return cost.get_item(c,record_id)
_revision_project_cache={}
def record_project(r):
 if r.get('project_id'):return r['project_id']
 if r.get('source_id') in ['START','CHANGE','CHANGE_NEW']:return 'PROJECT_A'
 if not str(r.get('record_id','')).startswith('REV:'):return None
 # A registered revision can originate from another site; verify its title.
 path=Path(r.get('source_path',''));expected=r.get('source_sha256')
 if not path.is_file() or sha(path)!=expected:return None
 if expected not in _revision_project_cache:
  from openpyxl import load_workbook
  workbook=load_workbook(path,read_only=True,data_only=True,keep_links=False);found=set()
  try:
   for sheet in workbook:
    for row in sheet.iter_rows(max_row=4):
     for cell in row:
      if not isinstance(cell.value,str):continue
      title=norm(cell.value)
      if '현장A' in title and ('공사' in title or '사업' in title):found.add('PROJECT_A')
      if '현장B' in title and ('공사' in title or '사업' in title):found.add('PROJECT_B')
  finally:workbook.close()
  _revision_project_cache[expected]=next(iter(found)) if len(found)==1 else None
 return _revision_project_cache[expected]
def compact(r):
 return {k:r.get(k) for k in ['record_id','name','specification','unit','section','source_id','scenario','source_path','source_sha256','source_sheet','source_row','status','project_id']}
def find_candidates(component,project,preset,discipline=None):
 ids=(preset.get('components',{}).get(component['component_id'],[]) if preset else [])
 if isinstance(ids,str):ids=[ids]
 found=[]
 for identifier in ids:
  r=get_record(identifier)
  if (record_project(r)==project['id'] and (not project.get('source_id') or r.get('source_id')==project['source_id']) and (project.get('scenario') is None or r.get('scenario')==project['scenario']) and unit(r['unit'])==unit(component['unit']) and (not discipline or r.get('department')==discipline) and r.get('status')=='VERIFIED_CONTRACT_RATE'):found.append(r)
 if found:return found
 terms=component.get('search_terms',[])
 if not terms:return []
 if str(project.get('source_id','')).startswith('REV:'):
  from working_catalog import HISTORY
  if not HISTORY.is_file():return []
  with sqlite3.connect(HISTORY.as_uri()+'?mode=ro',uri=True) as c:
   c.execute('PRAGMA query_only=ON')
   try:revision_rows=c.execute('SELECT payload_json FROM revision_catalog WHERE revision_id=?',(project['source_id'],)).fetchall()
   except sqlite3.OperationalError:return []
  for row in revision_rows:
   r=json.loads(row[0])
   if record_project(r)!=project['id'] or r.get('status')!='VERIFIED_CONTRACT_RATE':continue
   if project.get('scenario') is not None and r.get('scenario')!=project['scenario']:continue
   if discipline and r.get('department')!=discipline:continue
   if unit(r['unit'])==unit(component['unit']) and any(norm(t) in norm(' '.join(str(r.get(k,'')) for k in ['name','specification','section'])) for t in terms):found.append(r)
  return found[:100]
 if project['id']=='PROJECT_A':
  with cost.connect() as c:
   src=project.get('source_id')
   if not src:return []
   rows=c.execute("SELECT payload_json FROM app_catalog WHERE route='CONTRACT' AND json_extract(payload_json,'$.source_id')=?",(src,))
   for row in rows:
    r=json.loads(row[0]);expected=2 if r['department']=='토목' and src!='START' else 1
    if r.get('scenario')!=project.get('scenario',expected):continue
    if (r.get('status')=='VERIFIED_CONTRACT_RATE' and unit(r['unit'])==unit(component['unit']) and (not discipline or r.get('department')==discipline) and any(norm(t) in norm(' '.join(str(r.get(k,'')) for k in ['name','specification','section'])) for t in terms)):found.append(r)
 else:
  with rules() as c:
   for row in c.execute('SELECT payload_json FROM external_contracts WHERE project_id=?',(project['id'],)):
    r=json.loads(row[0])
    if project.get('source_id') and r.get('source_id')!=project['source_id']:continue
    if project.get('scenario') is not None and r.get('scenario')!=project['scenario']:continue
    if (r.get('status')=='VERIFIED_CONTRACT_RATE' and unit(r['unit'])==unit(component['unit']) and (not discipline or r.get('department')==discipline) and any(norm(t) in norm(' '.join(str(r.get(k,'')) for k in ['name','specification','section'])) for t in terms)):found.append(r)
 return found[:100]
def compose(req):
 if not isinstance(req,dict):raise AssemblyError('공사 요청 객체가 필요합니다')
 project=req.get('project') or {};work=req.get('work') or {}
 if not isinstance(project,dict) or not isinstance(work,dict):raise AssemblyError('project/work는 객체여야 합니다')
 if 'scenario' in project and (type(project['scenario']) is not int or project['scenario']<1):raise AssemblyError('scenario는 양의 정수여야 합니다')
 if not project.get('id'):raise AssemblyError('현장 project.id를 지정하세요')
 with rules() as c:
  if not c.execute('SELECT 1 FROM projects WHERE project_id=?',(project['id'],)).fetchone():raise AssemblyError('등록·검증된 현장 ID가 필요합니다')
  tid=work.get('type_id')
  if not tid:
   matches=list(c.execute('SELECT type_id FROM aliases WHERE alias=?',(norm(work.get('name','')),)))
   if len(matches)!=1:return {'status':'NEEDS_WORK_TYPE','request':req,'type_candidates':list_types(work.get('name','')),'missing':[{'reason':'공사 유형을 특정하세요'}],'lines':[]}
   tid=matches[0][0]
  row=c.execute('SELECT payload_json FROM assembly_types WHERE type_id=?',(tid,)).fetchone()
  if not row:raise AssemblyError('공사 유형 ID가 없습니다')
  kind=json.loads(row[0])
 presets=[p for p in kind.get('presets',[]) if p['project_id']==project['id'] and (not project.get('source_id') or p.get('source_id')==project['source_id'])]
 preset=presets[0] if len(presets)==1 else None
 if preset:
  project={**project,'source_id':project.get('source_id',preset.get('source_id')),'scenario':project.get('scenario',preset.get('scenario'))}
 conditions=work.get('conditions',{});selections=req.get('selections',{});missing=[];lines=[];excluded=[]
 if not isinstance(conditions,dict) or not isinstance(selections,dict):raise AssemblyError('조건/선택 형식을 확인하세요')
 if not project.get('source_id'):missing.append({'reason':'현장 내 계약/설계 원문 판본 source_id를 특정하세요'})
 if project.get('scenario') is None:missing.append({'reason':'총괄/연차 시나리오를 특정하세요'})
 component_ids={x['component_id'] for x in kind['components']}
 if set(selections)-component_ids:raise AssemblyError('알 수 없는 구성요소 선택: '+', '.join(sorted(set(selections)-component_ids)))
 for cond in kind.get('required_conditions',[]):
  value=conditions.get(cond['key'])
  if value is not None and type(value) is not bool:raise AssemblyError('조건 '+cond['key']+'는 실제 참/거짓이어야 합니다')
  if value is not True:missing.append({'component_id':None,'reason':cond.get('label',cond['key']),'condition':cond['key']})
 for component in kind['components']:
  cid=component['component_id'];condition=component.get('condition')
  active=True
  if condition:
   key=condition['key']
   if key not in conditions:
    missing.append({'component_id':cid,'reason':'적용 조건 미확정','condition':key});active=None
   else:
    expected=condition.get('equals',True);actual=conditions[key]
    if isinstance(expected,bool) and type(actual) is not bool:raise AssemblyError('조건 '+key+'는 실제 참/거짓이어야 합니다')
    active=actual==expected
  elif component.get('requirement')=='conditional':
   key='include_'+cid
   if key not in conditions:missing.append({'component_id':cid,'reason':'선택 공종 포함 여부 미확정','condition':key});active=None
   else:
    if type(conditions[key]) is not bool:raise AssemblyError(key+'는 참/거짓이어야 합니다')
    active=conditions[key]
  if active is False:
   if cid in selections:raise AssemblyError('제외한 구성요소에 단가 선택이 있습니다: '+cid)
   excluded.append({'component_id':cid,'reason':'명시 조건에 의해 제외','condition':condition});continue
  for cond in component.get('required_conditions',[]):
   if cond['key'] not in conditions:missing.append({'component_id':cid,'reason':cond.get('label',cond['key']),'condition':cond['key']})
  candidates=find_candidates(component,project,preset,kind.get('discipline'))
  selection=selections.get(cid)
  supplied=cid in selections
  selections_for_component=selection if isinstance(selection,list) else [selection]
  if supplied and (not selections_for_component or any(x is None for x in selections_for_component)):raise AssemblyError('공종 선택을 빈 배열/null로 생략할 수 없습니다')
  if len(selections_for_component)>100:raise AssemblyError('한 구성요소 선택은 100개까지 지원합니다')
  if not supplied:selections_for_component=[{'record_id':candidates[0]['record_id']}] if preset and len(candidates)==1 else [None]
  for i,pick in enumerate(selections_for_component,1):
   if isinstance(pick,str):pick={'record_id':pick}
   selected=None
   if pick is not None:
    if not isinstance(pick,dict) or not pick.get('record_id'):raise AssemblyError('선택 record_id가 필요합니다')
    selected=get_record(pick['record_id'])
    if record_project(selected)!=project['id']:raise AssemblyError('다른 현장의 계약 행을 선택할 수 없습니다')
    if project.get('source_id') and selected.get('source_id')!=project['source_id']:raise AssemblyError('다른 계약 판본을 선택할 수 없습니다')
    if project.get('scenario') is not None and selected.get('scenario')!=project['scenario']:raise AssemblyError('총괄/연차 구간이 다릅니다')
    if selected.get('status')!='VERIFIED_CONTRACT_RATE':raise AssemblyError('검증되지 않은 원문 행입니다')
    if unit(selected.get('unit'))!=unit(component['unit']):raise AssemblyError('구성요소와 선택 공종의 단위가 다릅니다')
    if selected.get('department')!=kind.get('discipline'):raise AssemblyError('선택 공종의 부서가 공사 유형과 다릅니다')
    if supplied and selected['record_id'] not in {x['record_id'] for x in candidates}:raise AssemblyError('구성요소 후보에 없는 공종입니다. 근거 매핑을 등록하세요')
    if not isinstance(pick.get('scope_key',''),str) or len(pick.get('scope_key',''))>200:raise AssemblyError('scope_key는 200자 이내 구간 문자열이어야 합니다')
   lid=cid if len(selections_for_component)==1 else cid+'#'+str(i)
   if selected is None:missing.append({'component_id':cid,'line_id':lid,'reason':'규격에 맞는 계약 공종 선택 필요','candidate_count':len(candidates)})
   lines.append({'line_id':lid,'component_id':cid,'name':component['name'],'role':component.get('role'),'unit':component['unit'],'active':active,'requirement':component.get('requirement'),'selected':compact(selected) if selected else None,'selection_basis':'EXPLICIT' if supplied else 'UNIQUE_PROJECT_PRESET' if selected else None,'candidates':[compact(x) for x in candidates],'quantity_rule':component.get('quantity_rule'),'rule_evidence':component.get('evidence',[]),'scope_note':component.get('scope_note'),'choice_group':component.get('choice_group'),'exclusive_with':component.get('exclusive_with',[]),'included_components':component.get('included_components',[]),'incompatible_selected_records':component.get('incompatible_selected_records',[]),'scope_key':pick.get('scope_key','') if pick else ''})
 active_ids={x['component_id'] for x in lines if x['active'] is True and x['selected']}
 required_any=kind.get('required_any_components',[])
 if required_any and not active_ids.intersection(required_any):missing.append({'reason':'공사에 최소 한 개의 주 공종을 선택해야 합니다','required_any_components':required_any})
 seen=set();groups={}
 for line in lines:
  if line['active'] is not True or not line['selected']:continue
  key=(line['selected']['record_id'],line['scope_key'])
  if key in seen:missing.append({'line_id':line['line_id'],'reason':'동일 계약 행·위치 중복 산입'})
  seen.add(key)
  for other in set(line['exclusive_with']+line['included_components']) & active_ids:
   missing.append({'line_id':line['line_id'],'reason':'복합비용 포함 또는 대체 공종 중복','other_component':other})
  for other in lines:
   if other['active'] is True and other.get('selected') and other['selected']['record_id'] in line.get('incompatible_selected_records',[]) and other.get('scope_key','')==line.get('scope_key',''):
    missing.append({'line_id':line['line_id'],'reason':'선택 기구에 이미 포함된 비용의 중복','other_component':other['component_id'],'included_record_id':other['selected']['record_id']})
  if line['choice_group']:groups.setdefault(line['choice_group'],set()).add(line['component_id'])
 for key,items in groups.items():
  if len(items)>1:missing.append({'reason':'대체 공법 동시 선택','choice_group':key,'components':sorted(items)})
 return {'schema_version':'construction-cost/composition/v1','status':'COMPOSED' if not missing else 'NEEDS_SELECTION_OR_CONDITIONS','type_id':kind['type_id'],'name':kind['name'],'project':project,'scope_confirmed':work.get('scope_confirmed') is True,'conditions':conditions,'lines':lines,'excluded':excluded,'missing':missing,'core_database_sha256':sha(cost.DB),'assembly_database_sha256':sha(RULE_DB),'scope_note':kind.get('scope_note'),'unsupported_methods_requiring_new_rules':kind.get('unsupported_methods_requiring_new_rules',[])}
def takeoff(req,composition=None,measurements=None):
 from quantity_engine import evaluate_quantity
 comp=composition or compose(req);rows=[];missing=list(comp.get('missing',[]))
 direct=req.get('quantities',{});dimensions=req.get('dimensions',{});per_line=req.get('component_inputs',{})
 if any(not isinstance(x,dict) for x in [direct,dimensions,per_line]):raise AssemblyError('물량·치수·공종별 입력은 객체여야 합니다')
 line_ids={x['line_id'] for x in comp.get('lines',[])}
 if (set(direct)|set(per_line))-line_ids:raise AssemblyError('활성 구성에 없는 물량/치수 행 ID가 있습니다')
 for line in comp.get('lines',[]):
  lid=line['line_id'];cid=line['component_id'];selected=line.get('selected')
  target=unit(selected['unit']) if selected else (line.get('quantity_rule') or {}).get('output_unit') or line.get('unit')
  if not target:
   rows.append({**line,'quantity':None,'unit':None,'quantity_status':'INCOMPLETE','quantity_evidence':{'missing':['공종 단위 미확정']}});continue
  rule=line.get('quantity_rule')
  if not isinstance(per_line.get(lid,{}),dict):raise AssemblyError('공종별 치수 입력은 객체여야 합니다')
  variables=(rule or {}).get('variables',{})
  if set(per_line.get(lid,{}))-set(variables):raise AssemblyError('공종 산식에 정의되지 않은 변수 입력: '+lid)
  inputs={k:v for k,v in {**dimensions,**per_line.get(lid,{})}.items() if k in variables}
  derived=evaluate_quantity(rule,inputs,measurements) if rule else None
  explicit=direct.get(lid)
  if explicit is not None:
   result=evaluate_quantity({'expression':'q','output_unit':target,'variables':{'q':{'unit':target}}},{'q':explicit},measurements)
   result['method']='USER_QUANTITY'
   if derived:result['derived_check']=derived
   if derived and derived.get('value') is not None and result.get('value') is not None:
    result['derived_comparison']={'value':derived['value'],'unit':derived['unit'],'difference':str(Decimal(result['value'])-Decimal(derived['value'])) if unit(derived['unit'])==unit(target) else None,'resolution':'명시 수량 우선; 도면/산식값 별도 보존'}
  elif derived is not None:result=derived;result['method']='FORMULA_OR_DRAWING'
  else:result={'status':'INCOMPLETE','value':None,'unit':target,'missing':['직접 물량 또는 검증된 수량 산식 필요']}
  if result.get('test_only') and req.get('purpose')!='TEST_ONLY':raise AssemblyError('시험용 근거는 TEST_ONLY 요청에서만 사용할 수 있습니다')
  if result.get('value') is not None and unit(result.get('unit'))!=unit(target):
   result={**result,'status':'BLOCKED','value':None,'errors':['산출 물량 단위와 선택 공종 단위가 다릅니다']}
  if line['active'] is not True:result={**result,'status':'INCOMPLETE','value':None,'missing':[*result.get('missing',[]),'공종 적용 조건 미확정']}
  if result.get('status')!='COMPLETE':missing.append({'line_id':lid,'reason':'물량 미확정','details':result})
  rows.append({**line,'quantity':result.get('value'),'unit':target,'quantity_status':result['status'],'quantity_evidence':result})
 return {'schema_version':'construction-cost/takeoff/v1','status':'TAKEOFF_COMPLETE' if not missing else 'INCOMPLETE','project':comp.get('project'),'type_id':comp.get('type_id'),'name':comp.get('name'),'scope_confirmed':comp.get('scope_confirmed',False),'composition':comp,'lines':rows,'missing':missing,'drawing_measurements':measurements or []}
def historical_amount(record,quantity):
 if record.get('status')!='VERIFIED_CONTRACT_RATE':raise AssemblyError('검증된 계약 행이 아닙니다')
 p=Path(record['source_path'])
 if not p.is_file() or sha(p)!=record['source_sha256']:raise AssemblyError('계약 원문 SHA가 달라졌습니다')
 q=cost.decimal(quantity,'수량');factor=cost.decimal(record.get('quantity_scale','1'))
 if q<0 and not record.get('signed_quantity_allowed'):raise AssemblyError('이 공종에 음수를 적용할 수 없습니다')
 if q>0 and record.get('signed_quantity_allowed'):raise AssemblyError('공제 공종에 양수 수량을 적용할 수 없습니다')
 parts={k:cost.decimal(record['unit_components'][k])*factor for k in ['material','labor','expense']}
 amounts={k:cost.trunc(q*v) for k,v in parts.items()}
 return {'status':'COMPLETE','record_id':record['record_id'],'quantity':str(q),'unit':record['unit'],'unit_components':{k:str(v) for k,v in parts.items()},'unit_total':str(sum(parts.values())),'amount_components':{k:str(v) for k,v in amounts.items()},'amount':str(sum(amounts.values())),'basis':record.get('price_basis','기존 계약 저장단가'),'source':compact(record),'source_label':record.get('source_label'),'price_date':record.get('price_date')}
def current_amount(record,q,selection,date,purpose,context=None):
 quantity=cost.decimal(q,'수량')
 if quantity<0 and not record.get('signed_quantity_allowed'):raise AssemblyError('일반 공종에 음수 수량을 적용할 수 없습니다')
 if quantity>0 and record.get('signed_quantity_allowed'):raise AssemblyError('공제 공종에 양수 수량을 적용할 수 없습니다')
 if not isinstance(selection,dict):return {'status':'INCOMPLETE','amount':None,'missing':['적용시점 가격 근거 미지정']}
 if selection.get('scope_confirmed') is not True or not selection.get('scope_statement'):return {'status':'INCOMPLETE','amount':None,'missing':['기존 계약과 신규 단가의 규격·포함범위 대조 필요']}
 if selection.get('contract_record_id')!=record['record_id']:raise AssemblyError('가격 비교 대상 계약 행 ID가 다릅니다')
 if selection.get('scope_evidence'):verify_file(selection['scope_evidence'])
 if selection.get('api_price_id'):
  if (selection.get('api_mapping') or {}).get('target_role')!='MATERIAL_SUPPLY' or any(word in record['name'] for word in ['설치','타설','부설','철거']):raise AssemblyError('API 직접단가는 별도 자재 공급행에만 적용합니다')
  from external_pricing import resolve_api_price
  price=resolve_api_price(selection,{**record,'kind':'material'},date,{**(context or {}),'purpose':purpose,'record_id':record['record_id'],'quantity':str(q),'unit':record['unit']})
  if price['historical']:raise AssemblyError('신규 가격안에 과거 저장가격이 포함됐습니다. 동일 결과 재생은 calculate에서 수행하세요')
  parts=price['cost_components'];amounts={k:cost.trunc(quantity*cost.decimal(v)) for k,v in parts.items()}
  return {**price,'status':'COMPLETE','amount':str(sum(amounts.values())),'unit_total':price['price'],'unit_components':parts,'amount_components':{k:str(v) for k,v in amounts.items()},'quantity':str(q),'unit':record['unit'],'scope_statement':selection['scope_statement']}
 if selection.get('catalog_record_id'):
  selected=get_record(selection['catalog_record_id'])
  if selected['route'] not in ['MARKET','PUMSEM']:raise AssemblyError('신규 산정안에 과거 계약행을 대체 적용할 수 없습니다')
  if unit(selected['unit'])!=unit(record['unit']):raise AssemblyError('비교 단위가 다릅니다')
  if (norm(selected['name'])!=norm(record['name']) or norm(selected.get('specification'))!=norm(record.get('specification'))) and not selection.get('scope_evidence'):raise AssemblyError('공표/품셈 규격과 기존 계약 범위의 차이에 대한 근거 매핑이 필요합니다')
  from pricing_policy import inherit_policy
  line={**inherit_policy(context or {},selection.get('calculation',{})),'record_id':selected['record_id'],'row_id':'CURRENT','quantity':str(q),'unit':selected['unit'],'pricing_date':date,'purpose':purpose,'price_data_mode':(context or {}).get('price_data_mode','CURRENT_REFRESH')}
  if line.get('allow_historical_prices'):raise AssemblyError('신규 산정안에 과거 관찰가격 적용을 허용하지 않습니다')
  with cost.connect() as c:out=cost.calculate_line(c,line)
  if any((x.get('price') or {}).get('historical') for x in out.get('resources',[])):raise AssemblyError('신규 가격안에 과거 자원가격이 포함됐습니다')
  return {**out,'status':'COMPLETE' if out['amount'] is not None else 'INCOMPLETE','scope_statement':selection['scope_statement'],'source_origin':'INDEPENDENT_EXTERNAL'}
 quote=selection.get('quote')
 if not isinstance(quote,dict):return {'status':'INCOMPLETE','amount':None,'missing':['공표 단가 또는 신규 견적 필요']}
 if quote.get('test_only') is True and purpose!='TEST_ONLY':raise AssemblyError('시험 견적은 실제 작업에 사용할 수 없습니다')
 for key in ['item_name','specification','unit','price','start_date','end_date','tax_basis','scope']:
  if quote.get(key) is None or str(quote[key]).strip()=='':raise AssemblyError('견적 근거 '+key+'가 필요합니다')
 source=verify_file(quote)
 from pricing_policy import validate_quote_origin,validate_quote_identity
 with cost.connect() as con:origin=validate_quote_origin(con,quote,context)
 if origin!='INDEPENDENT_EXTERNAL':raise AssemblyError('내역서 원문가격은 신규 독립가격안에 사용할 수 없습니다')
 validate_quote_identity(quote,record,selection,cost.canonical_unit)
 if quote.get('currency','KRW')!='KRW':raise AssemblyError('원화 단가가 필요합니다')
 start=datetime.date.fromisoformat(quote['start_date']);end=datetime.date.fromisoformat(quote['end_date'])
 if not start<=datetime.date.fromisoformat(date)<=end:raise AssemblyError('견적 유효기간 밖입니다')
 if unit(quote['unit'])!=unit(record['unit']):raise AssemblyError('견적 단위가 다릅니다')
 if (norm(quote['item_name'])!=norm(record['name']) or norm(quote['specification'])!=norm(record['specification'])) and not selection.get('quote_mapping'):raise AssemblyError('이름/규격 차이에 대한 명시 quote_mapping이 필요합니다')
 price=cost.decimal(quote['price'],'견적 단가');pack=cost.decimal(quote.get('pack_quantity','1'),'포장량')
 if price<0 or pack<=0:raise AssemblyError('음수 가격 또는 유효하지 않은 포장량')
 factor=Decimal(1)
 if quote['tax_basis']=='INCLUDED':
  vat=cost.decimal(quote.get('vat_percent'),'VAT 비율')
  if not 0<=vat<=100:raise AssemblyError('VAT 비율 범위 오류')
  factor/=1+vat/100
 elif quote['tax_basis']!='EXCLUDED':raise AssemblyError('세금 포함/별도 확인 필요')
 factor/=pack;price*=factor
 parts=quote.get('cost_components')
 if parts is not None:
  if not isinstance(parts,dict) or set(parts)!={'material','labor','expense'}:raise AssemblyError('견적 3비목을 정확히 명시하세요')
  values={k:cost.decimal(v) for k,v in parts.items()}
  if any(v<0 for v in values.values()) or sum(values.values())!=cost.decimal(quote['price']):raise AssemblyError('견적 원시 3비목 합계가 견적 가격과 다릅니다')
  values={k:v*factor for k,v in values.items()};amounts={k:cost.trunc(cost.decimal(q)*v) for k,v in values.items()}
  amount=sum(amounts.values());parts={k:str(v) for k,v in values.items()};amounts={k:str(v) for k,v in amounts.items()}
 else:amount=cost.trunc(cost.decimal(q)*price);amounts=None
 return {'status':'COMPLETE','amount':str(amount),'unit_total':str(price),'unit_components':parts,'amount_components':amounts,'quantity':str(q),'unit':quote['unit'],'basis':'TEST_ONLY 견적' if quote.get('test_only') else '기간·범위가 명시된 독립 견적','source':source,'quote':quote,'scope_statement':selection['scope_statement'],'source_origin':'INDEPENDENT_EXTERNAL'}
def compare(req,takeoff_result=None,measurements=None):
 from pricing_policy import validate_policy,reference_authorized
 validate_policy(req);allow_reference=reference_authorized(req)
 take=takeoff_result or takeoff(req,measurements=measurements);date=(req.get('project') or {}).get('pricing_date')
 if not isinstance(date,str):raise AssemblyError('실제 가격 적용일 project.pricing_date가 필요합니다')
 datetime.date.fromisoformat(date)
 rows=[];new_prices=req.get('new_prices',{})
 if not isinstance(new_prices,dict):raise AssemblyError('신규 가격 선택은 객체여야 합니다')
 if set(new_prices)-{x['line_id'] for x in take['lines']}:raise AssemblyError('비교 대상에 없는 신규 가격 행 ID가 있습니다')
 for line in take['lines']:
  contract={'status':'INCOMPLETE','amount':None};current={'status':'INCOMPLETE','amount':None};errors=[]
  if line['selected'] and line['quantity_status']=='COMPLETE':
   record=get_record(line['selected']['record_id'])
   if allow_reference:
    try:contract=historical_amount(record,line['quantity'])
    except (AssemblyError,cost.InputError,ValueError) as exc:errors.append(str(exc))
   else:contract={'status':'SUPPRESSED_BY_POLICY','record_id':record['record_id'],'unit':record['unit'],'amount':None,'basis':'원문 가격 참고 요청 없음','source':compact(record)}
   try:current=current_amount(record,line['quantity'],new_prices.get(line['line_id']),date,req.get('purpose'),req)
   except (AssemblyError,cost.InputError,ValueError) as exc:current={'status':'BLOCKED','amount':None,'missing':[str(exc)]}
  rows.append({'line_id':line['line_id'],'component_id':line['component_id'],'name':line['selected']['name'] if line['selected'] else line['name'],'specification':line['selected']['specification'] if line['selected'] else None,'quantity':line['quantity'],'unit':line['unit'],'contract':contract,'current':current,'difference':str(Decimal(current['amount'])-Decimal(contract['amount'])) if current['amount'] is not None and contract['amount'] is not None else None,'errors':errors,'quantity_evidence':line['quantity_evidence']})
 scope_ready=take['status']=='TAKEOFF_COMPLETE' and take['scope_confirmed'] is True
 totals={}
 for side in ['contract','current']:
  if side=='contract' and not allow_reference:
   totals[side]={'status':'SUPPRESSED_BY_POLICY','amount':None,'known_subtotal':None};continue
  known=sum((Decimal(x[side]['amount']) for x in rows if x[side].get('amount') is not None),Decimal(0))
  done=scope_ready and bool(rows) and all(x[side].get('amount') is not None for x in rows)
  totals[side]={'status':'COMPLETE' if done else 'INCOMPLETE','amount':str(known) if done else None,'known_subtotal':str(known)}
 complete=all(x['status']=='COMPLETE' for x in totals.values())
 status='COMPARISON_COMPLETE' if complete else ('EXTERNAL_SCOPE_COMPLETE' if not allow_reference and totals['current']['status']=='COMPLETE' else 'PARTIAL_COMPARISON')
 return {'schema_version':'construction-cost/price-comparison/v1','status':status,'pricing_policy':'EXTERNAL_PRIMARY','contract_reference':req.get('contract_reference'),'price_data_mode':req.get('price_data_mode','CURRENT_REFRESH'),'purpose':req.get('purpose','PROJECT_WORK'),'project':take['project'],'pricing_date':date,'type_id':take['type_id'],'scope':'선택 공사 구성의 직접 구성금액. 제경비·세금·전체 도급액은 별도 연결 검증 대상.','scope_confirmed':take['scope_confirmed'],'takeoff_status':take['status'],'missing':take['missing']+([] if take['scope_confirmed'] else [{'reason':'공사 전체 포함·제외 범위 확인 필요'}]),'rows':rows,'totals':totals,'difference_total':str(Decimal(totals['current']['amount'])-Decimal(totals['contract']['amount'])) if complete else None,'core_database_sha256':sha(cost.DB),'assembly_database_sha256':sha(RULE_DB)}
